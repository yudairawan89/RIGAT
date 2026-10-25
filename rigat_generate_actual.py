#!/usr/bin/env python3
"""RIGAT measured-results analysis and Excel-driven table/figure generation.

No artificial predictions, reference scores, or cost scenarios are used here.
"""
from __future__ import annotations
import argparse
import hashlib
import importlib.metadata
import json
import math
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats
from scipy.optimize import minimize
from scipy.special import expit, logit
from sklearn.isotonic import IsotonicRegression
from sklearn.metrics import (accuracy_score, average_precision_score, confusion_matrix,
    f1_score, matthews_corrcoef, precision_recall_curve, roc_auc_score, roc_curve)

VERSION = "2.0.2"
CALIBRATORS = ["Raw A4", "Temperature Scaling", "Platt", "Isotonic", "Beta",
               "Entropy-Only", "Distance-Only", "Adaptive Trust (A5)"]
METHOD_LABELS = {"A0": "A0 DeBERTa Prefix", "A1": "A1 Hierarchical Mean",
    "A2": "A2 Reliability", "A2_DRO": "A2 + GroupDRO, no invariance",
    "A3": "A3 Invariance", "A4": "A4 Worst Environment", "A5": "A5 Full RIGAT",
    "DP-Net-adapted": "DP-Net-adapted", "WaveDetect-adapted": "WaveDetect-adapted",
    "Longformer": "Longformer", "Binoculars": "Binoculars", "A4_Distance_Only": "A4 Distance-Only"}


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def json_safe(value):
    """Use JSON null, never the nonstandard NaN token, for unavailable fields."""
    if isinstance(value, dict):
        return {k: json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    if isinstance(value, np.generic):
        return json_safe(value.item())
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value

def raw_entropy(logits):
    """Actual normalized binary entropy, always in [0, 1]."""
    p = expit(np.asarray(logits, dtype=float))
    p = np.clip(p, 1e-15, 1 - 1e-15)
    return -(p * np.log(p) + (1 - p) * np.log1p(-p)) / np.log(2.0)

def probability(logits):
    return np.clip(expit(logits), 1e-12, 1 - 1e-12)

def nll_value(y, logits):
    return float(np.mean(np.logaddexp(0.0, logits) - y * logits))

def risk_curve(y, logits):
    """AURC with expected error order within exact confidence ties.

    This avoids class-order artifacts for isotonic plateaus. Absolute logits
    preserve ranking without probability saturation for temperature scaling.
    """
    conf_score = np.abs(logits)
    errors = ((logits >= 0).astype(int) != y).astype(float)
    order = np.argsort(-conf_score, kind="stable")
    scores, errors = conf_score[order], errors[order]
    starts = np.r_[0, np.flatnonzero(np.diff(scores) != 0) + 1]
    ends = np.r_[starts[1:], len(scores)]
    totals = np.add.reduceat(errors, starts)
    expected = np.repeat(totals / (ends - starts), ends - starts)
    risks = np.cumsum(expected) / np.arange(1, len(y) + 1)
    grid = np.linspace(0.01, 1.0, 100)
    idx = np.maximum(0, np.ceil(grid * len(y)).astype(int) - 1)
    return float(risks.mean()), grid, risks[idx]

def metrics(y, logits, n_bins=15, curves=True):
    p = probability(logits)
    pred = (logits >= 0).astype(int)
    hits = (pred == y).astype(float)
    confidence = np.maximum(p, 1 - p)
    edges = np.linspace(0.5, 1.0, n_bins + 1)
    bin_index = np.minimum(np.searchsorted(edges, confidence, side="right") - 1, n_bins - 1)
    rows, ece = [], 0.0
    for b in range(n_bins):
        mask = bin_index == b
        n = int(mask.sum())
        conf = float(confidence[mask].mean()) if n else np.nan
        acc = float(hits[mask].mean()) if n else np.nan
        if n:
            ece += n / len(y) * abs(acc - conf)
        rows.append({"Bin": b + 1, "Bin_Low": edges[b], "Bin_High": edges[b + 1],
                     "N_bin": n, "Mean_Confidence": conf, "Empirical_Accuracy": acc})
    cm = confusion_matrix(y, pred, labels=[0, 1])
    auroc = roc_auc_score(1 - y, -logits) if len(np.unique(y)) == 2 else np.nan
    auprc = average_precision_score(1 - y, -logits) if (y == 0).any() else np.nan
    aurc, grid, risk = risk_curve(y, logits) if curves else (np.nan, None, None)
    result = {"Accuracy": float(accuracy_score(y, pred)),
              "Macro_F1": float(f1_score(y, pred, labels=[0, 1], average="macro", zero_division=0)),
              "MCC": float(matthews_corrcoef(y, pred)), "AUROC": float(auroc), "AUPRC": float(auprc),
              "TP_Machine": int(cm[0, 0]), "FN_Machine": int(cm[0, 1]),
              "FP_Human": int(cm[1, 0]), "TN_Human": int(cm[1, 1]),
              "Machine_Recall": float(cm[0, 0] / max(1, cm[0].sum())),
              "Human_Recall": float(cm[1, 1] / max(1, cm[1].sum())),
              "Human_FPR": float(cm[1, 0] / max(1, cm[1].sum())),
              "ECE": float(ece), "NLL": nll_value(y, logits),
              "Brier": float(np.mean((p - y) ** 2)), "AURC": aurc}
    return result, rows, grid, risk

@dataclass
class Calibrator:
    name: str
    parameters: dict = field(default_factory=dict)
    iso: object = None
    fit_info: dict = field(default_factory=dict)

    def transform(self, logits, entropy, distance):
        p = self.parameters
        if self.name == "Raw A4":
            return np.asarray(logits).copy()
        if self.name == "Temperature Scaling":
            return logits / p["T"]
        if self.name == "Platt":
            return p["Slope"] * logits + p["Intercept"]
        if self.name == "Beta":
            q = probability(logits)
            return p["Beta_A"] * np.log(q) - p["Beta_B"] * np.log1p(-q) + p["Intercept"]
        if self.name == "Isotonic":
            return logit(np.clip(self.iso.predict(expit(logits)), 1e-12, 1 - 1e-12))
        tau = expit(p["Bias"] - p["Alpha"] * entropy - p["Beta"] * distance)
        return np.maximum(tau, 1e-12) / p["T"] * logits

    def serializable(self):
        obj = {"name": self.name, "parameters": self.parameters, "fit_info": self.fit_info}
        if self.iso is not None:
            obj["X_thresholds"] = self.iso.X_thresholds_.tolist()
            obj["y_thresholds"] = self.iso.y_thresholds_.tolist()
        return obj

def fit_calibrator(name, y, logits, entropy, distance, maxiter=600, trust_l2=1e-4):
    start = time.perf_counter()
    if name == "Raw A4":
        return Calibrator(name, fit_info={"Success": True, "Status": "Identity", "Iterations": 0,
                                         "Function_Evaluations": 0, "Fit_CPU_Sec": 0.0})
    if name == "Isotonic":
        iso = IsotonicRegression(y_min=1e-12, y_max=1 - 1e-12, out_of_bounds="clip")
        iso.fit(expit(logits), y)
        fitted_mse = float(np.mean((iso.predict(expit(logits)) - y) ** 2))
        return Calibrator(name, {"Threshold_Count": int(len(iso.X_thresholds_))}, iso,
                          {"Success": True, "Status": "Fitted (PAVA)", "Iterations": np.nan,
                           "Function_Evaluations": np.nan, "Gradient_Inf_Norm": np.nan,
                           "Fit_CPU_Sec": time.perf_counter() - start, "Optimizer": "PAVA",
                           "Objective": "MSE (isotonic least squares)", "Lambda_L2": 0.0,
                           "Regularization_Penalty": 0.0,
                           "Optimizer_Objective_Value": fitted_mse})
    def decode(theta):
        if name == "Temperature Scaling":
            return {"T": float(np.exp(theta[0]))}
        if name == "Platt":
            return {"Slope": float(np.exp(theta[0])), "Intercept": float(theta[1])}
        if name == "Beta":
            return {"Beta_A": float(np.exp(theta[0])), "Beta_B": float(np.exp(theta[1])),
                    "Intercept": float(theta[2])}
        return {"T": float(np.exp(theta[0])),
                "Alpha": 0.0 if name == "Distance-Only" else float(theta[2]),
                "Beta": 0.0 if name == "Entropy-Only" else float(theta[-1]),
                "Bias": float(theta[1])}
    if name == "Temperature Scaling":
        starts, bounds = [[0.0]], [(-3.0, 3.0)]
    elif name == "Platt":
        starts, bounds = [[0.0, 0.0]], [(-4.0, 4.0), (-8.0, 8.0)]
    elif name == "Beta":
        starts, bounds = [[0.0, 0.0, 0.0]], [(-4.0, 4.0), (-4.0, 4.0), (-8.0, 8.0)]
    else:
        dimensions = 4 if name == "Adaptive Trust (A5)" else 3
        starts = [[0.0, 3.0] + [0.5] * (dimensions - 2),
                  [0.3, 4.0] + [1.0] * (dimensions - 2),
                  [math.log(0.9), logit(0.9)] + [0.0] * (dimensions - 2)]
        bounds = [(-3.0, 3.0), (-8.0, 12.0)] + [(0.0, 8.0)] * (dimensions - 2)
    brier_weight = 0.5 if name in ["Entropy-Only", "Distance-Only", "Adaptive Trust (A5)"] else 0.0
    ridge = trust_l2 if brier_weight else 0.0
    def objective(theta):
        cal = Calibrator(name, decode(theta))
        transformed = cal.transform(logits, entropy, distance)
        p = expit(transformed)
        data_loss = nll_value(y, transformed) + brier_weight * float(np.mean((p - y) ** 2))
        # T and gate bias can trade off when the sigmoid approaches its exponential
        # tail. A weak, explicit ridge stabilizes that scale/bias degeneracy.
        # It changes the objective, not the reported NLL/ECE/Brier observations.
        return data_loss + ridge * float(np.dot(theta, theta))
    results = [minimize(objective, x0, method="L-BFGS-B", bounds=bounds,
                        options={"maxiter": maxiter, "ftol": 1e-12, "gtol": 1e-7, "maxls": 40})
               for x0 in starts]
    finite = [r for r in results if np.isfinite(r.fun)]
    if not finite:
        raise RuntimeError(f"No finite fit for {name}")
    best = min(finite, key=lambda r: r.fun)
    grad = np.asarray(best.jac, dtype=float)
    projected = grad.copy()
    bound_count = 0
    for i, (lo, hi) in enumerate(bounds):
        if best.x[i] <= lo + 1e-6:
            bound_count += 1
            projected[i] = min(0.0, grad[i])
        elif best.x[i] >= hi - 1e-6:
            bound_count += 1
            projected[i] = max(0.0, grad[i])
    info = {"Success": bool(best.success), "Status": str(best.message), "Optimizer": "L-BFGS-B",
            "Iterations": int(best.nit), "Function_Evaluations": int(sum(r.nfev for r in results)),
            "Gradient_Inf_Norm": float(np.max(np.abs(grad))),
            "Projected_Gradient_Inf_Norm": float(np.max(np.abs(projected))),
            "Parameters_At_Bounds": bound_count, "Fit_CPU_Sec": time.perf_counter() - start,
            "Objective": "NLL + 0.5*Brier + lambda_L2*||theta||^2" if brier_weight else "NLL",
            "Lambda_L2": ridge, "Regularization_Penalty": ridge * float(np.dot(best.x, best.x)),
            "Optimizer_Objective_Value": float(best.fun), "Starts": len(starts)}
    return Calibrator(name, decode(best.x), fit_info=info)

def calibration_diagnosis(seed, cals, instance):
    y, logits, u, d = instance
    before = metrics(y, logits, curves=False)[0]
    rows = []
    for name, cal in cals.items():
        if name == "Raw A4":
            continue
        after = metrics(y, cal.transform(logits, u, d), curves=False)[0]
        weight = 0.5 if name in ["Entropy-Only", "Distance-Only", "Adaptive Trust (A5)"] else 0.0
        data_before = before["Brier"] if name == "Isotonic" else before["NLL"] + weight * before["Brier"]
        data_after = after["Brier"] if name == "Isotonic" else after["NLL"] + weight * after["Brier"]
        objective_metric = "MSE" if name == "Isotonic" else "NLL + 0.5*Brier" if weight else "NLL"
        rows.append({"Seed": seed, "Calibration_Partition": "valid_calibration",
                     "N_Validation": len(y), "Calibrator": name,
                     "Val_NLL_Before": before["NLL"], "Val_NLL_After": after["NLL"],
                     "NLL_Reduction": before["NLL"] - after["NLL"],
                     "Val_ECE_Before": before["ECE"], "Val_ECE_After": after["ECE"],
                     "Val_Brier_Before": before["Brier"], "Val_Brier_After": after["Brier"],
                     "Data_Objective_Metric": objective_metric,
                     "Data_Objective_Before": data_before,
                     "Data_Objective_After": data_after,
                     **cal.parameters, **cal.fit_info})
    return rows

def runtime_arguments(argv=None):
    """Strip only Jupyter's connection-file option from implicit kernel argv.

    Explicit main([...]) arguments stay strict, as do normal terminal arguments.
    Do not use parse_known_args: it would silently hide misspelled user options.
    """
    if argv is not None:
        return list(argv)
    arguments = list(sys.argv[1:])
    launcher = Path(sys.argv[0]).name.lower() if sys.argv else ""
    in_kernel = "ipykernel" in sys.modules or launcher in {"ipykernel_launcher.py", "ipykernel_launcher"}
    if not in_kernel:
        return arguments
    filtered = []
    i = 0
    while i < len(arguments):
        arg = arguments[i]
        if (arg in {"-f", "--f"} and i + 1 < len(arguments)
                and arguments[i + 1].lower().endswith(".json")):
            i += 2
            continue
        if arg.startswith(("-f=", "--f=")) and arg.split("=", 1)[1].lower().endswith(".json"):
            i += 1
            continue
        filtered.append(arg)
        i += 1
    return filtered


def write_workbook(path, frames, readme):
    """Preserve the agreed Python Excel workflow; README is exactly one row."""
    from openpyxl.styles import Alignment, Font, PatternFill
    path = Path(path)
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
        pd.DataFrame([["Data provenance", readme]]).to_excel(writer, sheet_name="README", header=False, index=False)
        for name, frame in frames.items():
            if len(frame) > 1048575:
                raise ValueError(f"Excel row limit exceeded: {name}; keep per-instance predictions in sidecar files")
            frame.to_excel(writer, sheet_name=name[:31], index=False)
        for sheet in writer.book:
            sheet.sheet_view.showGridLines = False
            if sheet.title == "README":
                sheet.column_dimensions["A"].width = 18
                sheet.column_dimensions["B"].width = 130
                sheet.row_dimensions[1].height = min(409, max(140, 16*math.ceil(len(readme)/105)))
                sheet["A1"].font = Font(bold=True, color="FFFFFF")
                sheet["A1"].fill = PatternFill("solid", fgColor="172A3A")
                for c in sheet[1]:
                    c.alignment = Alignment(wrap_text=True, vertical="center")
                continue
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 44
            for c in sheet[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="172A3A")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for column in sheet.columns:
                header = str(column[0].value)
                width = min(42, max(17, len(header)+2))
                if header in {"Method", "Calibrator", "Protocol", "Condition", "Comparator", "Method_Label"}:
                    width = max(32, width)
                if header in {"File", "Checkpoint", "SHA256", "Checkpoint_SHA256"}:
                    width = 70
                if header in {"Value", "Configuration"}:
                    width = 110
                sheet.column_dimensions[column[0].column_letter].width = width
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.00E+00" if header.startswith("p_") else "0.000000"
                    cell.alignment = Alignment(vertical="center", wrap_text=isinstance(cell.value, str))
                    if isinstance(cell.value, str):
                        sheet.row_dimensions[cell.row].height = min(409, max(sheet.row_dimensions[cell.row].height or 15,
                            15*math.ceil(len(cell.value)/(width*.9))))
    temporary.replace(path)


def export_training_workbook(directory):
    directory = Path(directory)
    config = json.loads((directory / "experiment_config.json").read_text(encoding="utf-8"))
    collections, registry = {}, []
    for complete in sorted((directory / "runs").glob("*/run_complete.json")):
        run = json.loads(complete.read_text(encoding="utf-8"))
        registry.append({k: v for k, v in run.items() if k != "records"})
        for name, record in run["records"].items():
            path = directory / record["path"]
            if sha256(path) != record["sha256"]:
                raise ValueError(f"Record hash mismatch: {path}")
            collections.setdefault(name, []).append(pd.read_csv(path))
    frames = {name: pd.concat(parts, ignore_index=True) for name, parts in collections.items()}
    frames["Run_Registry"] = pd.DataFrame(registry)
    frames["Dataset_Composition"] = pd.DataFrame(config["partition_composition"])
    configuration_rows = []
    def flatten(value, prefix):
        if isinstance(value, dict):
            for key, child in value.items():
                flatten(child, f"{prefix}.{key}" if prefix else str(key))
        elif isinstance(value, list):
            for index, child in enumerate(value):
                flatten(child, f"{prefix}[{index}]")
        else:
            configuration_rows.append({"Component": prefix, "Value": value})
    flatten(config, "")
    frames["Configuration"] = pd.DataFrame(configuration_rows)
    fits = len(frames.get("Calibration_Diagnosis", []))
    coverage_fits = len(frames.get("Coverage_Fits", []))
    baseline_fits = len(frames.get("Baseline_Fits", []))
    runs = sum(x["Independent_Training_Run_Count"] for x in registry)
    completed = {(x["Method"], x["Seed"]) for x in registry if x["Independent_Training_Run_Count"]}
    pending = [{"Method": method, "Seed": seed} for method in config["methods"] for seed in config["seeds"] if (method, seed) not in completed]
    binoculars_complete = any(x["Method"] == "Binoculars" for x in registry)
    status = {"training_runs_planned": len(config["methods"])*len(config["seeds"]),
        "training_runs_completed": runs, "pending_stage_i": pending,
        "main_calibrator_fits_completed": fits, "coverage_calibrator_fits_completed": coverage_fits,
        "baseline_scale_fits_completed": baseline_fits,
        "binoculars_required": config.get("with_binoculars", False), "binoculars_completed": binoculars_complete,
        "requested_runs_complete": not pending and (not config.get("with_binoculars") or binoculars_complete),
        "external_condition_enabled": "External" in config["counts"], "mode": config["mode"]}
    (directory / "completion_summary.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
    readme = (
        f"Results originate from the executed training/evaluation pipeline in {config['mode']} mode, "
        f"using backbone {config['backbone']}; {runs} Stage I training runs completed; {fits} main calibrator fits "
        f"and {coverage_fits} additional coverage fits; {baseline_fits} separate baseline scale fits. A5 shares the A4 Stage I checkpoint and adds no training. "
        "Per-document predictions, source hashes, checkpoints and configuration accompany the workbook; only completed runs enter the summaries. "
        "Training time includes token-cache retrieval, input preparation, forward/backward passes and optimizer updates; "
        "epoch validation, checkpoint I/O, initial tokenization and export are excluded. End_to_End_ms includes raw-text tokenization, "
        "input preparation/transfer, the model and probability mapping; A5 also includes distance and trust. CUDA_Model_ms "
        "measures the model/aggregation using CUDA Events, excluding tokenization, transfer and CPU calibration. CPU runs leave GPU measurements unavailable. "
        "Peak VRAM reports the PyTorch allocated/reserved allocator, rather than total device usage. Fit_CPU_Sec is elapsed "
        "wall-clock fitting time on the CPU, separate from reference/input encoding. Distance uses per-seed training statistics "
        f"and Ledoit-Wolf estimation on class residuals; distance normalization={config.get('distance_normalization', 'train-median')}. Preprocessing details and "
        "interpretation limits are documented in TRAINING_GUIDE.md. External datasets and comparators are reported only when evaluated "
        "or when corresponding actual predictions are available. Binoculars is evaluated once; rows are repeated across the seed schedule only for "
        "paired statistics, not independent training. DP-Net-adapted uses a discrete Gaussian REINFORCE policy; "
        "WaveDetect-adapted applies CWT to token-probability signals from the DeBERTa encoder. Adaptation details are documented "
        "in the guide and do not claim replication of the complete original implementations. No performance targets or forced differences between seeds are imposed.")
    write_workbook(directory / "RIGAT_Training_Results.xlsx", frames, readme)
    (directory / "README_results.md").write_text(readme, encoding="utf-8")
    return directory / "RIGAT_Training_Results.xlsx"


def aggregate(frame, keys, columns):
    if frame.empty:
        return pd.DataFrame(columns=keys+[x+s for x in columns for s in ["_Mean", "_SD"]]+["n_runs"])
    functions = {"n_runs": ("Seed", "nunique")}
    for col in columns:
        functions[col+"_Mean"] = (col, "mean")
        functions[col+"_SD"] = (col, "std")
    return frame.groupby(keys, as_index=False, sort=False).agg(**functions)


def compare_actual(classification, calibration):
    rows = []
    def add(family, condition, metric, comparator, delta, positive):
        delta = np.asarray(delta, dtype=float)
        n = len(delta)
        mean = float(np.mean(delta))
        sd = float(np.std(delta, ddof=1)) if n > 1 else np.nan
        lo, hi, p, statistic = np.nan, np.nan, np.nan, np.nan
        if n >= 2:
            if np.all(np.abs(delta) < 1e-14):
                mean, sd, lo, hi, p, statistic = 0., 0., 0., 0., 1., 0.
            elif sd > 1e-14:
                statistic, p = stats.ttest_1samp(delta, 0)
                half = stats.t.ppf(.975, n-1)*sd/np.sqrt(n)
                lo, hi = mean-half, mean+half
            else:
                lo = hi = mean
        rows.append({"Family": family, "Condition": condition, "Metric": metric, "Comparator": comparator,
            "n_runs": n, "Paired_Delta_Mean": mean, "Paired_Delta_SD": sd, "CI_95_Low": lo,
            "CI_95_High": hi, "t_statistic": statistic, "p_raw": p,
            "Preferred_Delta": "positive" if positive else "negative", "df": n-1,
            "Uncertainty_Unit": "Aligned seeds; fixed comparator" if comparator == "Binoculars" else "Aligned seeds; t-interval"})
    for condition, frame in classification.groupby("Condition", sort=False):
        reference = frame[frame.Method == "A5"].set_index("Seed")
        external = condition == "External"
        for metric in (["Macro_F1", "AUROC", "ECE", "NLL", "Brier", "AURC"] if external else ["Macro_F1", "AUROC", "AUPRC", "Human_FPR"]):
            for method, other in frame[frame.Method != "A5"].groupby("Method", sort=False):
                pairs = reference[[metric]].join(other.set_index("Seed")[[metric]], lsuffix="_a", rsuffix="_b").dropna()
                if len(pairs):
                    family = ("External classification" if metric in ["Macro_F1", "AUROC"] else "External calibration") if external else "Classification"
                    add(family, condition, metric, method, pairs[metric+"_a"]-pairs[metric+"_b"], metric in ["Macro_F1", "AUROC", "AUPRC"])
        if not external:
            pivot = frame.pivot(index="Seed", columns="Method", values="Macro_F1")
            if set(["A2", "A2_DRO", "A3", "A4"]) <= set(pivot.columns):
                delta = (pivot.A4-pivot.A3-pivot.A2_DRO+pivot.A2).dropna()
                if len(delta):
                    add("Factorial interaction", condition, "Macro_F1", "Invariance x GroupDRO", delta, True)
    if not calibration.empty:
        for condition, frame in calibration[calibration.Condition != "External"].groupby("Condition", sort=False):
            reference = frame[frame.Calibrator == "Adaptive Trust (A5)"].set_index("Seed")
            for metric in ["ECE", "NLL", "Brier", "AURC"]:
                for name, other in frame[frame.Calibrator != "Adaptive Trust (A5)"].groupby("Calibrator", sort=False):
                    pairs = reference[[metric]].join(other.set_index("Seed")[[metric]], lsuffix="_a", rsuffix="_b").dropna()
                    if len(pairs):
                        add("Calibration", condition, metric, name, pairs[metric+"_a"]-pairs[metric+"_b"], False)
    if not rows:
        return pd.DataFrame(columns=["Family", "Condition", "Metric", "Comparator", "n_runs", "Paired_Delta_Mean", "p_raw", "p_Holm"])
    result = pd.DataFrame(rows)
    result["p_Holm"] = np.nan
    for _, indices in result.groupby(["Family", "Condition", "Metric"]).groups.items():
        ordered = result.loc[indices].sort_values("p_raw", na_position="last").index
        adjusted = 0.0
        for j, index in enumerate(ordered):
            p = result.loc[index, "p_raw"]
            if pd.notna(p):
                adjusted = max(adjusted, min(1., (len(ordered)-j)*p))
                result.loc[index, "p_Holm"] = adjusted
    return result


def timing_summary(frame, keys):
    if frame.empty:
        return pd.DataFrame()
    return frame.groupby(keys, as_index=False).agg(
        Timing_Runs=("Timing_ID", "nunique"), End_to_End_p50_ms=("End_to_End_ms", "median"),
        End_to_End_p95_ms=("End_to_End_ms", lambda x: x.quantile(.95)),
        CUDA_Model_p50_ms=("CUDA_Model_ms", "median"),
        CUDA_Model_p95_ms=("CUDA_Model_ms", lambda x: x.quantile(.95)))


def build_tables(frames):
    cls = frames.get("Classification", pd.DataFrame())
    cal = frames.get("Calibration", pd.DataFrame())
    if cls.empty:
        raise ValueError("Workbook has no completed classification evaluations")
    standard = cls[cls.Condition == "Standard"]
    desired_order = ["A0", "Longformer", "Binoculars", "DP-Net-adapted", "WaveDetect-adapted", "A1", "A2", "A2_DRO", "A3", "A4", "A5"]
    order = [m for m in desired_order if m in set(standard.Method)]
    order += [m for m in standard.Method.unique() if m not in order]
    composition = frames["Dataset_Composition"].rename(columns={"Split": "Dataset_Split"})
    configuration = frames["Configuration"].rename(columns={"Value": "Configuration"})
    def setting(name, default):
        selected = frames["Configuration"]
        selected = selected[selected.Component == name]
        if selected.empty:
            return default
        value = selected.iloc[0].Value
        try:
            return json.loads(value) if isinstance(value, str) else value
        except (ValueError, TypeError):
            return value
    chunk_length = int(setting("chunk_length", 512))
    max_chunks = int(setting("max_chunks", 8))
    t01 = pd.DataFrame([{"Method": m, "Role": "Reference comparator" if m in {"Longformer", "Binoculars", "DP-Net-adapted", "WaveDetect-adapted"} else "Component comparison",
        "Evaluation_Conditions": "Standard; OOD GPT; Paraphrase", "Calibration_Comparison": m in {"A4", "A5"}} for m in order])
    t04 = pd.DataFrame([{"Method": m, "Reliability": m in ["A2", "A2_DRO", "A3", "A4", "A5"],
        "Invariance": m in ["A3", "A4", "A5"], "GroupDRO": m in ["A2_DRO", "A4", "A5"],
        "Document_Coverage": f"{int(setting('longformer_length', 4096))}-token prefix" if m == "Longformer" else f"{int(setting('binoculars_max_tokens', 512))}-token scoring window" if m == "Binoculars" else f"{chunk_length}-token prefix" if m in {"A0", "DP-Net-adapted"} else f"Up to {max_chunks} document-wide {chunk_length}-token chunks",
        "Calibration": "Adaptive Trust" if m == "A5" else "Development threshold and positive scale" if m == "Binoculars" else "None",
        "n_runs": int(standard[standard.Method == m].Seed.nunique())} for m in order])
    t05 = aggregate(standard, ["Method"], ["Macro_F1", "Accuracy", "MCC", "AUROC", "AUPRC", "Human_FPR"])
    t05 = t05.set_index("Method").reindex(order).reset_index()
    wide = cls[cls.Condition != "External"].pivot(index=["Seed", "Method"], columns="Condition", values="Macro_F1").reset_index()
    robust_columns = []
    for condition in ["Standard", "OOD GPT", "Paraphrase"]:
        if condition in wide:
            robust_columns.append(condition)
    for condition, target in [("OOD GPT", "OOD_Drop"), ("Paraphrase", "Paraphrase_Drop")]:
        if "Standard" in wide and condition in wide:
            wide[target] = wide.Standard-wide[condition]
            robust_columns.append(target)
    t06 = aggregate(wide, ["Method"], robust_columns)
    t06 = t06.set_index("Method").reindex(order).reset_index()
    t07 = aggregate(cal[cal.Condition != "External"] if not cal.empty else cal, ["Condition", "Calibrator"], ["ECE", "NLL", "Brier", "AURC", "Accuracy", "Macro_F1", "AUROC", "AUPRC"])
    if not cal.empty:
        for metric in ["ECE", "NLL", "Brier", "AURC"]:
            values = cal.pivot(index=["Seed", "Condition"], columns="Calibrator", values=metric)
            delta = (values["Adaptive Trust (A5)"]-values["Distance-Only"]).groupby("Condition").mean()
            t07["A5_minus_Distance_"+metric] = np.where(t07.Calibrator == "Adaptive Trust (A5)", t07.Condition.map(delta), np.nan)
    coverage = frames.get("Coverage", pd.DataFrame())
    coverage_timing = timing_summary(frames.get("Coverage_Timing", pd.DataFrame()), ["Seed", "Protocol", "Kmax"])
    if not coverage.empty and not coverage_timing.empty:
        coverage = coverage.merge(coverage_timing, on=["Seed", "Protocol", "Kmax"], validate="one_to_one")
    coverage_columns = [c for c in ["Macro_F1", "ECE", "Token_Coverage_Mean", "Complete_Token_Document_Fraction",
        "Observed_Token_Chunks_Mean", "Word_Coverage_Mean", "Complete_Word_Document_Fraction", "Observed_512Word_Chunks_Mean",
        "End_to_End_p50_ms", "End_to_End_p95_ms", "CUDA_Model_p50_ms", "CUDA_Model_p95_ms"] if c in coverage]
    t08 = aggregate(coverage, ["Protocol", "Kmax"], coverage_columns)
    external = cls[cls.Condition == "External"]
    t11 = aggregate(external, ["Method"], ["Macro_F1", "AUROC", "ECE", "NLL", "Brier", "AURC"])
    for table in [t05, t06, t11]:
        if "Method" in table and "Binoculars" in set(table.Method):
            table.loc[table.Method == "Binoculars", "n_runs"] = 1
            table.loc[table.Method == "Binoculars", table.columns.str.endswith("_SD")] = np.nan
    if "Binoculars" in set(t04.Method):
        t04.loc[t04.Method == "Binoculars", "n_runs"] = 1
    cost = frames.get("Cost", pd.DataFrame()).copy()
    timed = timing_summary(frames.get("Timing", pd.DataFrame()), ["Seed", "Method"])
    if not cost.empty and not timed.empty:
        cost = cost.merge(timed, on=["Seed", "Method"], validate="one_to_one")
    cost_columns = [c for c in cost if c not in {"Method", "Seed", "Execution_Device", "Training_Cost_Parent", "Parent_Stage_I"}
                    and pd.api.types.is_numeric_dtype(cost[c])]
    t12 = aggregate(cost, ["Method"], cost_columns)
    if not cost.empty:
        counts = cost.groupby("Method").Independent_Training_Run_Count.sum()
        t12["Completed_Independent_Stage_I_Runs"] = t12.Method.map(counts).astype(int)
        t12["Training_Cost_Parent"] = t12.Method.map(cost.groupby("Method").Training_Cost_Parent.first())
        for field in ["Execution_Device", "Precision_Executed", "Hardware_CPU", "Hardware_GPU"]:
            if field in cost:
                t12[field] = t12.Method.map(cost.groupby("Method")[field].first())
    tables = {"T01_Comparator_Overview": t01, "T02_Dataset_Composition": composition,
        "T03_Experimental_Configuration": configuration, "T04_Ablation_and_Baselines": t04,
        "T05_Standard_Test": t05, "T06_Robustness_Shifts": t06, "T07_Calibration_Trust": t07,
        "T08_Coverage_Efficiency": t08, "T09_Statistical_Comparisons": compare_actual(cls, cal),
        "T10_Calibration_Diagnosis": frames.get("Calibration_Diagnosis", pd.DataFrame()),
        "T11_RAID_Validation": t11, "T12_Computational_Cost": t12}
    for table in tables.values():
        if "Method" in table:
            table.insert(table.columns.get_loc("Method")+1, "Method_Label", table.Method.map(METHOD_LABELS).fillna(table.Method))
    comparison = tables["T09_Statistical_Comparisons"]
    comparison.insert(comparison.columns.get_loc("Comparator")+1, "Comparator_Label", comparison.Comparator.map(METHOD_LABELS).fillna(comparison.Comparator))
    return tables


def verify_frames(frames):
    cls = frames["Classification"]
    if cls.duplicated(["Seed", "Method", "Condition"]).any():
        raise ValueError("Duplicate classification keys")
    for _, row in cls.iterrows():
        if int(row.TP_Machine+row.FN_Machine+row.FP_Human+row.TN_Human) != row.N:
            raise ValueError("Confusion matrix does not sum to population")
    if "A5" in set(cls.Method):
        a4 = cls[cls.Method == "A4"].set_index(["Seed", "Condition"])
        a5 = cls[cls.Method == "A5"].set_index(["Seed", "Condition"])
        cells = ["TP_Machine", "FN_Machine", "FP_Human", "TN_Human", "Macro_F1"]
        if not np.allclose(a4.loc[a5.index, cells], a5[cells], atol=1e-12):
            raise ValueError("A4/A5 decisions differ at fixed threshold")
    if "Domain" in frames:
        domain = frames["Domain"].groupby(["Seed", "Method", "Condition"])[["TP_Machine", "FN_Machine", "FP_Human", "TN_Human"]].sum()
        overall = cls.set_index(["Seed", "Method", "Condition"])[domain.columns]
        if not domain.sort_index().equals(overall.sort_index()):
            if not np.allclose(domain.sort_index(), overall.sort_index()):
                raise ValueError("Domain totals differ from overall counts")
    if "Risk" in frames:
        risk = frames["Risk"]
        endpoint = risk[risk.Coverage == 1].set_index(["Seed", "Condition", "Calibrator"])
        cal = frames["Calibration"].set_index(["Seed", "Condition", "Calibrator"])
        if not np.allclose(endpoint.Selective_Risk.sort_index(), (1-cal.Accuracy).sort_index(), atol=1e-12):
            raise ValueError("Risk-coverage endpoint differs from classification error")
    return True


def render_figures(out, frames, tables, dpi):
    import matplotlib
    from textwrap import fill
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Serif", "font.size": 10,
        "axes.spines.top": False, "axes.spines.right": False, "savefig.facecolor": "white"})
    colors = plt.get_cmap("tab10").colors
    cls = frames["Classification"]
    standard = tables["T05_Standard_Test"]
    methods = standard.Method.tolist()
    label = lambda m: METHOD_LABELS.get(m, m)
    def save(fig, name):
        fig.savefig(out / (name+".png"), dpi=dpi, bbox_inches="tight")
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(12, 5.6), layout="constrained")
    ablation = standard[standard.Method.str.startswith("A")]
    ax.bar(range(len(ablation)), ablation.Macro_F1_Mean, yerr=ablation.Macro_F1_SD.fillna(0), capsize=4, color=colors[0])
    ax.set(xticks=range(len(ablation)), xticklabels=[fill(label(x), width=18) for x in ablation.Method],
        ylabel="Macro-F1 (mean ± SD)", title="Standard-test component comparison", ylim=(0, 1))
    save(fig, "Fig3_Ablation_MacroF1")
    target = "A5" if "A5" in methods else methods[-1]
    seed = int(cls.Seed.min())
    conditions = [c for c in ["Standard", "OOD GPT", "Paraphrase"] if c in set(cls.Condition)]
    fig, axes = plt.subplots(1, len(conditions), figsize=(5*len(conditions), 4.5), layout="constrained", squeeze=False)
    for ax, condition in zip(axes[0], conditions):
        r = cls[(cls.Method == target) & (cls.Seed == seed) & (cls.Condition == condition)].iloc[0]
        cm = np.array([[r.TP_Machine, r.FN_Machine], [r.FP_Human, r.TN_Human]], dtype=int)
        norm = cm/np.maximum(1, cm.sum(1, keepdims=True))
        ax.imshow(norm, vmin=0, vmax=1, cmap="viridis")
        for i in range(2):
            for j in range(2):
                ax.text(j, i, f"{cm[i,j]:,}\n({norm[i,j]:.3f})", ha="center", va="center", color="black" if norm[i,j]>.5 else "white")
        ax.set(xticks=[0, 1], yticks=[0, 1], xticklabels=["Machine", "Human"], yticklabels=["Machine", "Human"],
            title=f"{condition}\n{target}, seed {seed}", xlabel="Predicted", ylabel="Actual")
    save(fig, "Fig4_Confusion_Matrices")
    curves = frames.get("ROC_PR", pd.DataFrame())
    if not curves.empty:
        fig, axes = plt.subplots(1, 2, figsize=(13, 5), layout="constrained")
        for method, group in curves[(curves.Condition == "Standard") & (curves.Seed == seed)].groupby("Method", sort=False):
            if method not in {"A0", "A1", "A2", "A3", "A4", "A5"}:
                continue
            axes[0].plot(group.Grid, group.ROC_TPR, label=label(method))
            axes[1].plot(group.Grid, group.PR_Precision, label=label(method))
        axes[0].set(xlabel="Human false positive rate", ylabel="Machine recall", title=f"ROC, seed {seed}")
        axes[1].set(xlabel="Machine recall", ylabel="Machine precision", title=f"Precision–recall, seed {seed}")
        for ax in axes:
            ax.legend(fontsize=8)
            ax.grid(alpha=.2)
        save(fig, "Fig5_ROC_PR_Curves")
    fig, axes = plt.subplots(1, 2, figsize=(15, 7), layout="constrained")
    for method, group in cls[cls.Condition.isin(conditions)].groupby("Method", sort=False):
        if method not in {"A5", "WaveDetect-adapted", "DP-Net-adapted", "Binoculars", "Longformer", "A0"}:
            continue
        g = group.groupby("Condition").Macro_F1.agg(["mean", "std"]).reindex(conditions)
        axes[0].errorbar(range(len(conditions)), g["mean"], yerr=g["std"].fillna(0), marker="o", capsize=3, label=label(method))
    axes[0].set(xticks=range(len(conditions)), xticklabels=conditions, ylabel="Macro-F1 (mean ± SD)")
    axes[0].legend(fontsize=8)
    domains = frames["Domain"]
    pivot = domains[domains.Method == target].pivot_table(index="Domain", columns="Condition", values="Macro_F1").reindex(columns=conditions)
    cmap = plt.get_cmap("viridis").copy()
    cmap.set_bad("#eeeeee")
    im = axes[1].imshow(np.ma.masked_invalid(pivot.to_numpy()), aspect="auto", cmap=cmap, vmin=0, vmax=1)
    axes[1].set(xticks=range(len(conditions)), xticklabels=conditions, yticks=range(len(pivot)), yticklabels=pivot.index, title=f"{target}: domain Macro-F1")
    for i in range(len(pivot)):
        for j in range(len(conditions)):
            v = pivot.iloc[i, j]
            axes[1].text(j, i, "—" if pd.isna(v) else f"{v:.3f}", ha="center", va="center", color="black" if pd.isna(v) or v>.5 else "white", fontsize=8)
    fig.colorbar(im, ax=axes[1], shrink=.7)
    save(fig, "Fig6_Robustness_and_Domain")
    cal = tables["T07_Calibration_Trust"]
    if not cal.empty:
        names = list(dict.fromkeys(cal.Calibrator))
        fig, axes = plt.subplots(1, 2, figsize=(15, 6), layout="constrained")
        width = .8/len(names)
        for i, name in enumerate(names):
            g = cal[cal.Calibrator == name].set_index("Condition").reindex(conditions)
            for ax, metric in zip(axes, ["ECE", "AURC"]):
                ax.bar(np.arange(len(conditions))+(i-(len(names)-1)/2)*width, g[metric+"_Mean"], width,
                       yerr=g[metric+"_SD"].fillna(0), capsize=2, color=colors[i%10], label=name)
                ax.set(xticks=range(len(conditions)), xticklabels=conditions, ylabel=metric)
        handles, labels = axes[0].get_legend_handles_labels()
        fig.legend(handles, labels, loc="outside lower center", ncol=4, fontsize=8)
        save(fig, "Fig7_ECE_AURC_Bar")
    co = tables["T08_Coverage_Efficiency"]
    if not co.empty:
        fig, axes = plt.subplots(1, 3, figsize=(19, 5.5), layout="constrained")
        for i, (protocol, group) in enumerate(co.groupby("Protocol", sort=False)):
            g = group.sort_values("Kmax")
            if protocol.startswith("Fixed K="):
                axes[0].errorbar(g.End_to_End_p95_ms_Mean, g.Macro_F1_Mean, yerr=g.Macro_F1_SD.fillna(0),
                    xerr=g.End_to_End_p95_ms_SD.fillna(0), marker="o", capsize=3, label=protocol)
                for r in g.itertuples():
                    axes[0].annotate(f"K={r.Kmax}", (r.End_to_End_p95_ms_Mean, r.Macro_F1_Mean),
                        xytext=(5, 7), textcoords="offset points", fontsize=8)
            axes[1].errorbar(g.Kmax, g.ECE_Mean, yerr=g.ECE_SD.fillna(0), marker="o", capsize=3, label=protocol)
        g = co[co.Protocol.str.startswith("Fixed K=")].sort_values("Kmax")
        for column, text in [("Word_Coverage_Mean_Mean", "Descriptive word coverage"),
                             ("Complete_Word_Document_Fraction_Mean", "Complete documents (words)"),
                             ("Token_Coverage_Mean_Mean", "Encoder token coverage"),
                             ("Complete_Token_Document_Fraction_Mean", "Complete documents (tokens)")]:
            axes[2].plot(g.Kmax, 100*g[column], marker="o", label=text)
        axes[0].set(xlabel="Measured end-to-end p95 (ms/document)", ylabel="Macro-F1")
        axes[1].set(xlabel="Maximum token chunks K", ylabel="ECE")
        axes[2].set(xlabel="K", ylabel="Coverage (%)", ylim=(0, 101))
        for ax in axes:
            ax.legend(fontsize=7)
            ax.grid(alpha=.2)
        save(fig, "Fig8_Latency_Coverage")
    selected_names = ["Raw A4", "Temperature Scaling", "Distance-Only", "Adaptive Trust (A5)"]
    if "Bins" in frames:
        fig, axes = plt.subplots(1, len(conditions), figsize=(5.5*len(conditions), 5.3), layout="constrained", squeeze=False)
        for ax, condition in zip(axes[0], conditions):
            ax.plot([.5, 1], [.5, 1], "k--", lw=.8)
            for name in selected_names:
                sub = frames["Bins"]
                sub = sub[(sub.Condition == condition) & (sub.Calibrator == name) & (sub.N_bin > 0)]
                rows = [(np.average(g.Mean_Confidence, weights=g.N_bin), np.average(g.Empirical_Accuracy, weights=g.N_bin),
                         g.Empirical_Accuracy.std(ddof=1) if len(g)>1 else 0.) for _, g in sub.groupby("Bin")]
                if rows:
                    x, y, deviation = zip(*rows)
                    ax.errorbar(x, y, yerr=np.nan_to_num(deviation), capsize=2, marker="o", markersize=3, label=name)
            ax.set(xlim=(.49, 1.01), ylim=(-.02, 1.02), xlabel="Mean confidence", ylabel="Bin accuracy", title=condition)
            ax.legend(fontsize=7)
        save(fig, "Fig9_Reliability_Diagrams")
    if "Risk" in frames:
        fig, axes = plt.subplots(1, len(conditions), figsize=(5.5*len(conditions), 5.3), layout="constrained", squeeze=False)
        for ax, condition in zip(axes[0], conditions):
            for name in selected_names:
                sub = frames["Risk"]
                group = sub[(sub.Condition == condition) & (sub.Calibrator == name)].groupby("Coverage").Selective_Risk.agg(["mean", "std"])
                ax.plot(group.index, group["mean"], label=name)
                ax.fill_between(group.index, np.maximum(0, group["mean"]-group["std"].fillna(0)),
                                group["mean"]+group["std"].fillna(0), alpha=.1)
            ax.set(xlabel="Retained fraction", ylabel="Selective error", title=condition)
            ax.legend(fontsize=7)
        save(fig, "Fig10_Risk_Coverage_Curves")
    external = tables["T11_RAID_Validation"]
    if not external.empty:
        fig, axes = plt.subplots(1, 2, figsize=(16, 6), layout="constrained")
        external_labels = {"A0":"A0\nDeBERTa", "A4":"A4\nUncalibrated",
            "A4_Distance_Only":"A4\nDistance-only", "A5":"A5\nRIGAT",
            "DP-Net-adapted":"DP-Net\nadapted", "WaveDetect-adapted":"WaveDetect\nadapted"}
        for ax, metric in zip(axes, ["Macro_F1", "ECE"]):
            ax.bar(range(len(external)), external[metric+"_Mean"], yerr=external[metric+"_SD"].fillna(0), capsize=3)
            ax.set(xticks=range(len(external)), xticklabels=[external_labels.get(x,x) for x in external.Method],
                ylabel=metric, title="External evaluation")
            ax.tick_params(axis="x", labelsize=8)
        save(fig, "Fig11_RAID_Transfer")
    costs = tables["T12_Computational_Cost"].set_index("Method")
    fig, axes = plt.subplots(1, 2, figsize=(15, 6), layout="constrained")
    selected = [m for m in ["A0", "Longformer", "DP-Net-adapted", "WaveDetect-adapted", "A5"] if m in set(costs.index)]
    for i, method in enumerate(selected):
        score = standard[standard.Method == method].iloc[0]
        axes[0].errorbar(costs.loc[method, "Total_Params_M_Mean"], score.Macro_F1_Mean,
            yerr=0 if pd.isna(score.Macro_F1_SD) else score.Macro_F1_SD, fmt="o", capsize=3,
            color=colors[i], label=label(method))
    axes[0].set(xlabel="Total model parameters (millions)", ylabel="Standard Macro-F1", title="Model size and performance")
    axes[0].legend(fontsize=8)
    x = np.arange(len(selected))
    axes[1].bar(x, costs.reindex(selected).Additional_Stage_I_Wall_Sec_Mean/3600, color=colors[1])
    axes[1].set(xticks=x, xticklabels=[m.replace("-adapted", "\nadapted") for m in selected])
    for ax in axes:
        ax.grid(alpha=.2)
    axes[1].set(ylabel="Additional Stage I training wall time (hours)", title="Training cost, A5 shares A4")
    save(fig, "Fig12_Model_Size_vs_F1")


def main(argv=None):
    p = argparse.ArgumentParser(description="Generate RIGAT tables/figures from measured training Excel only")
    p.add_argument("--workbook", type=Path, required=True)
    p.add_argument("--output-dir", type=Path, default=Path("RIGAT_Actual_Report"))
    p.add_argument("--dpi", type=int, default=600)
    p.add_argument("--skip-figures", action="store_true")
    p.add_argument("--source-figures", type=Path,
                   help="Copy existing figures when preserving supplied T01-T12 tables")
    args = p.parse_args(runtime_arguments(argv))
    if args.dpi < 72:
        p.error("dpi must be >=72")
    workbook = pd.ExcelFile(args.workbook)
    from rigat_workbook_io import TABLE_NAMES, export_preserved_report
    if all(name in workbook.sheet_names for name in TABLE_NAMES):
        workbook.close()
        return export_preserved_report(args.workbook, args.output_dir,
            skip_figures=args.skip_figures, source_figures=args.source_figures, version=VERSION)
    readme = pd.read_excel(workbook, sheet_name="README", header=None).iloc[0, 1]
    frames = {name: pd.read_excel(workbook, sheet_name=name) for name in workbook.sheet_names if name != "README"}
    workbook.close()
    verify_frames(frames)
    tables = build_tables(frames)
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    for name, frame in tables.items():
        frame.to_csv(out / (name+".csv"), index=False, encoding="utf-8-sig")
    write_workbook(out / "RIGAT_Actual_Report.xlsx", tables, readme)
    if not args.skip_figures:
        render_figures(out, frames, tables, args.dpi)
    metadata = {"source_workbook": str(args.workbook.resolve()), "sha256": sha256(args.workbook),
                "version": VERSION, "tables": list(tables), "dpi": args.dpi,
                "comparisons": len(tables["T09_Statistical_Comparisons"]),
                "external_evaluated": not tables["T11_RAID_Validation"].empty,
                "source_script_sha256": sha256(__file__), "python_version": sys.version,
                "packages": {name: importlib.metadata.version(name) for name in ["numpy", "pandas", "scipy", "scikit-learn", "openpyxl", "matplotlib"]}}
    (out / "report_manifest.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(f"Report complete: {out}; {metadata['comparisons']} paired comparisons", flush=True)
    return out


if __name__ == "__main__":
    main()
